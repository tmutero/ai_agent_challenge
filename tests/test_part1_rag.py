"""Part 1: ingestion, routing, aggregation and fact verification -- offline (fake embedder, no LLM)."""

from dataclasses import replace

import pytest

from part1_rag.aggregate import run_aggregation
from part1_rag.answerer import Fact, GroundedAnswer, Passage, unsupported_numbers, verify_facts
from part1_rag.engine import RagEngine, default_settings
from part1_rag.loaders import load_knowledge_base
from part1_rag.query_transform import Aggregation, FieldFilter, QueryPlan, validate_plan


@pytest.fixture(scope="module")
def kb():
    return load_knowledge_base()


@pytest.fixture(scope="module")
def engine(kb):
    return RagEngine(replace(default_settings(), embedding_provider="fake"), kb)


def test_one_document_per_record(kb):
    by_type = {}
    for d in kb.documents:
        by_type[d.doc_type] = by_type.get(d.doc_type, 0) + 1
    assert by_type == {"csv_row": 20, "json_record": 4, "pdf_table_row": 13, "pdf_text": 2, "catalog": 6}
    # The PDF's second table continues across a page break without a header row.
    assert len(kb.datasets["investment_policy_t2"].frame) == 8
    # Nested JSON lists become a child table that cites the parent record.
    tiers = kb.datasets["fee_schedule_admin_fee_tiers"].frame
    assert len(tiers) == 10 and tiers["_chunk_id"].iloc[0] == "fee_schedule.json#UT::0"


def test_records_carry_provenance(kb):
    doc = next(d for d in kb.documents if d.doc_id == "funds.csv#KST-ESG12")
    assert doc.metadata["record_ref"] == "row 12"
    assert "Return 5-year %: not reported" in doc.text
    assert doc.metadata["fields"]["asset_class"] == "SA Equity"


def _plan(**kw):
    base = dict(standalone_query="q", intent="lookup", datasets=[], filters=[], sub_queries=[],
                aggregation=None, reasoning="")
    return QueryPlan(**{**base, **kw})


def test_validate_plan_drops_hallucinated_filters(kb):
    plan = _plan(datasets=["funds", "made_up"], filters=[
        FieldFilter(dataset="funds", column="risk_rating", values=["Very High", "High"]),
        FieldFilter(dataset="funds", column="ter_pct", values=["1.0"])])
    cleaned, notes = validate_plan(plan, kb)
    assert cleaned.datasets == ["funds"]
    assert [f.values for f in cleaned.filters] == [["High"]]
    assert len(notes) == 3


def test_routes_scope_filters_to_their_dataset(engine, kb):
    plan = _plan(datasets=["funds", "investment_policy_t1"], filters=[
        FieldFilter(dataset="investment_policy_t1", column="Client risk profile", values=["Conservative"])])
    routes = engine._routes(plan)
    assert {"category": "funds"} in routes
    assert {"category": "investment_policy_t1", "metadata.fields.client_risk_profile": "Conservative"} in routes
    assert {"category": "investment_policy_text"} in routes        # sibling prose of the same PDF
    chunks = engine._retrieve_all(["conservative"], routes)
    suit = [c for c in chunks if c.chunk.metadata.get("category") == "investment_policy_t1"]
    assert [c.chunk.doc_id for c in suit] == ["investment_policy.pdf#t1-Conservative"]


def test_aggregation_min_with_filter_and_missing_values(kb):
    spec = Aggregation(dataset="funds", operation="min", column="ter_pct", limit=3, return_columns=["fund_name"],
                       filters=[FieldFilter(dataset="funds", column="asset_class", values=["SA Equity"])])
    res = run_aggregation(spec, kb)
    assert res.rows["fund_code"].tolist() == ["KST-DIV07", "KST-EQ06", "KST-ESG12"]
    assert res.value == 1.14
    spec5 = Aggregation(dataset="funds", operation="max", column="return_5y_pct", limit=1, return_columns=[], filters=[])
    res5 = run_aggregation(spec5, kb)
    assert res5.rows["fund_code"].tolist() == ["KST-GEQ09"]
    assert set(res5.excluded_missing) == {"KST-INF11", "KST-ESG12"}


def test_fact_verification():
    passages = [Passage("funds.csv - row 3", "Total expense ratio (TER) %: 0.55\nMinimum lump sum (ZAR): 10000")]
    ans = GroundedAnswer(answer="", insufficient_evidence=False, facts=[
        Fact(claim="TER", value="0.55%", passage=1),
        Fact(claim="min", value="R10,000", passage=1),
        Fact(claim="wrong", value="0.65", passage=1),
        Fact(claim="bad cite", value="0.55", passage=2)])
    assert [c.verified for c in verify_facts(ans, passages)] == [True, True, False, False]


def test_unsupported_numbers_in_prose():
    passages = [Passage("p", "TER %: 1.14\nTER %: 1.37")]
    assert unsupported_numbers("Lowest is 1.14% [1], then 1.37% [2], average 1.26%.", passages) == ["1.26"]


# --- uploads: arbitrary files through the generic loaders ----------------------------------

from part1_rag.loaders import SourceFile, build_knowledge_base  # noqa: E402


def test_unknown_csv_infers_key_title_and_filters():
    csv = b"sku,name,colour,price_usd\nA1,Desk lamp,black,25\nA2,Desk chair,grey,120\nA3,Monitor arm,black,\n"
    kb, [rep] = build_knowledge_base([SourceFile("catalogue.csv", csv)])
    ds = kb.datasets["catalogue"]
    assert rep.ok and rep.records == 3
    assert ds.key_column == "sku" and ds.filterable["colour"] == ["black", "grey"]
    doc = next(d for d in kb.documents if d.doc_id == "catalogue.csv#A3")
    assert "row 3: Monitor arm" in doc.text and "Price (USD): not reported" in doc.text


def test_json_top_level_list_and_excel_and_text(tmp_path):
    import docx
    import pandas as pd
    xlsx = tmp_path / "t.xlsx"
    with pd.ExcelWriter(xlsx) as w:
        pd.DataFrame({"region": ["N", "S"], "sales": [10, 20]}).to_excel(w, sheet_name="Q1", index=False)
        pd.DataFrame({"region": ["N", "S"], "sales": [30, 40]}).to_excel(w, sheet_name="Q2", index=False)
    d = docx.Document()
    d.add_heading("Leave policy", 1)
    d.add_paragraph("Employees receive 21 days of annual leave per year, accrued monthly.")
    t = d.add_table(rows=2, cols=2)
    t.rows[0].cells[0].text, t.rows[0].cells[1].text = "Grade", "Days"
    t.rows[1].cells[0].text, t.rows[1].cells[1].text = "Senior", "25"
    docx_path = tmp_path / "policy.docx"
    d.save(docx_path)
    files = [SourceFile("people.json", b'[{"id": 1, "name": "Ann", "tags": ["a", "b"]}, {"id": 2, "name": "Bo"}]'),
             SourceFile("sales.xlsx", xlsx.read_bytes()),
             SourceFile("policy.docx", docx_path.read_bytes()),
             SourceFile("notes.md", b"# Intro\n\nThis knowledge base describes the quarterly sales process in detail.\n")]
    kb, reports = build_knowledge_base(files)
    assert all(r.ok for r in reports), reports
    assert {"people", "sales_q1", "sales_q2", "policy_t1"} <= set(kb.datasets)
    assert {"policy_text", "notes_text"} <= set(kb.text_sources)
    para = next(d for d in kb.documents if d.source == "policy.docx" and d.doc_type == "docx_text")
    assert "Leave policy" in para.metadata["record_ref"]


def test_bad_and_unsupported_files_are_reported_not_fatal():
    kb, reports = build_knowledge_base([SourceFile("x.exe", b"MZ"), SourceFile("broken.json", b"{not json"),
                                        SourceFile("ok.csv", b"a,b\n1,2\n")])
    by = {r.source: r for r in reports}
    assert "unsupported" in by["x.exe"].error and "could not parse" in by["broken.json"].error
    assert by["ok.csv"].ok and kb.sources == ["ok.csv"]


def test_empty_knowledge_base_answers_without_llm():
    kb, _ = build_knowledge_base([])
    r = RagEngine(replace(default_settings(), embedding_provider="fake"), kb).answer("anything")
    assert r.refused and "No documents" in r.refused
